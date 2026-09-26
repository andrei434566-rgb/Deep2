from __future__ import annotations

import csv
import json
import math
import re
import unicodedata
from dataclasses import replace
from pathlib import Path
from typing import Any, Iterable

from .depth import format_depth, meters_to_centimeters, normalize_depth
from .models import ColumnMapping, DescriptionRow, Issue


SUPPORTED_TABLES = {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls", ".csv", ".tsv"}

ROLE_ALIASES: dict[str, tuple[str, ...]] = {
    "well": ("скваж", "скв", "well", "borehole", "hole"),
    "interval": ("интервал фаци", "интервал слоя", "facies interval", "depth interval", "интервал", "глубин"),
    "top": ("кровл", "верх", "начало", "от", "from", "top", "start"),
    "base": ("подошв", "низ", "конец", "до", "to", "base", "bottom", "end"),
    "facies_thickness": ("толщина фаци", "мощность фаци", "толщина слоя", "мощность слоя", "facies thickness"),
    "core_top": ("отбор керна кровл", "интервал керна кровл", "core top", "core from"),
    "core_base": ("отбор керна подошв", "интервал керна подошв", "core base", "core to"),
    "label": ("название фаци", "наименование фаци", "литофаци", "класс", "метка", "label", "class", "facies name", "порода", "литология"),
    "class_code": ("код фаци", "фациальн код", "class code", "facies code", "код"),
    "class_index": ("индекс фаци", "фациальн индекс", "class index", "facies index", "индекс"),
    "description": ("литологическое описание", "краткое описание", "описан", "характерист", "description", "comment", "примечан"),
    "target_text": ("краткое описание", "стандартизированное описание", "target text", "target description"),
    "association": ("ассоциация фаций", "facies association"),
    "environment": ("обстановка осадконакопления", "depositional environment", "environment"),
    "field_name": ("месторожд", "площадь", "field", "site", "object"),
}


def normalize_text(value: Any) -> str:
    text = " ".join(str(value or "").replace("\n", " ").split())
    return unicodedata.normalize("NFKC", text).casefold().replace("ё", "е")


def display_text(value: Any) -> str:
    return " ".join(str(value or "").replace("\n", " ").split())


def well_key(value: str) -> str:
    text = re.sub(r"\b(?:скважина|скв|well|borehole)\b", "", normalize_text(value))
    return re.sub(r"[^a-zа-я0-9]+", "", text)


def as_float(value: Any) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else None
    text = re.sub(r"[\s\u00a0\u202f]+", "", str(value))
    match = re.search(r"-?\d+(?:[.,]\d+)?", text)
    parsed = float(match.group().replace(",", ".")) if match else None
    return parsed if parsed is not None and math.isfinite(parsed) else None


def parse_interval(value: Any) -> tuple[float | None, float | None]:
    text = display_text(value).replace("−", "-").replace("–", "-").replace("—", "-")
    values = re.findall(r"\d+(?:[.,]\d+)?", text)
    if len(values) < 2:
        return None, None
    return as_float(values[0]), as_float(values[1])


def _score(header: str, role: str) -> int:
    text = normalize_text(header)
    if not text:
        return 0
    score = 0
    for alias in ROLE_ALIASES[role]:
        alias_normalized = normalize_text(alias)
        if text == alias_normalized:
            score = max(score, 180 + len(alias_normalized))
        elif len(alias_normalized) <= 3 and re.search(rf"(?<!\w){re.escape(alias_normalized)}(?!\w)", text):
            score = max(score, 100 + len(alias_normalized))
        elif len(alias_normalized) > 3 and alias_normalized in text:
            score = max(score, 100 + len(alias_normalized))
    is_core = any(word in text for word in ("керн", "core"))
    is_facies = any(word in text for word in ("фаци", "слоя", "facies", "layer"))
    is_drilling = any(word in text for word in ("по бурен", "буров", "drilling"))
    is_gis = any(word in text for word in ("по гис", "gis", "logging"))
    is_top = any(word in text for word in ("кровл", "верх", "начало", " from", " top", " start"))
    is_base = any(word in text for word in ("подошв", "низ", "конец", " to", " base", " bottom", " end"))
    if role == "core_top" and is_core and is_top:
        score = max(score, 240)
    if role == "core_base" and is_core and is_base:
        score = max(score, 240)
    if role == "top" and is_facies and is_top:
        score = max(score, 230)
    if role == "base" and is_facies and is_base:
        score = max(score, 230)
    if role == "interval" and is_facies:
        score = max(score, 170)
    if role == "facies_thickness":
        if any(word in text for word in ("толщина фаци", "мощность фаци", "facies thickness")):
            score = max(score, 300)
        elif any(word in text for word in ("толщина слоя", "мощность слоя")):
            score = max(score, 220)
    if score <= 0:
        return 0
    if role in {"core_top", "core_base"}:
        score += 80 if is_core else -100
    if role in {"top", "base", "interval"}:
        score += 40 if is_facies else 0
        score -= 80 if is_core else 0
        # Training masks use the facies interval measured by drilling. A GIS
        # interval can have the same subheaders and must not win by position.
        score += 120 if is_facies and is_drilling else 0
        score -= 70 if is_facies and is_gis and not is_drilling else 0
    if role == "target_text" and "краткое описание" in text:
        # Source workbooks may place this field at column 22, 30 or anywhere
        # else. The semantic header is the contract, not the column number.
        score = max(score, 340)
    if role == "label" and any(word in text for word in ("код", "индекс", "code", "index")):
        score -= 120
    if role == "class_code" and "индекс" in text:
        score -= 100
    if role == "class_index" and "код" in text:
        score -= 100
    return max(0, score)


def _best(headers: list[str], role: str, used: set[int]) -> int | None:
    ranked = sorted(
        ((index + 1, _score(value, role)) for index, value in enumerate(headers) if index + 1 not in used),
        key=lambda item: (item[1], -item[0]),
        reverse=True,
    )
    return ranked[0][0] if ranked and ranked[0][1] > 0 else None


def _best_facies_pair(
    headers: list[str], rows: list[list[Any]], header_row: int,
    used: set[int], thickness_column: int | None,
    interval_source: str | None = None,
) -> tuple[int | None, int | None]:
    """Choose drilling facies limits together and verify them by facies thickness.

    Wide field forms contain several almost identical pairs named ``Кровля`` /
    ``Подошва``. Header meaning is primary, while the row-level equality
    ``Подошва - Кровля == Толщина фации`` is an independent safeguard against
    accidentally selecting the much wider core-sampling interval.
    """
    def belongs_to_source(header: str) -> bool:
        text = normalize_text(header)
        is_facies = any(word in text for word in ("фаци", "слоя", "facies", "layer"))
        is_drilling = any(word in text for word in ("по бурен", "буров", "drilling"))
        is_gis = any(word in text for word in ("по гис", "gis", "logging"))
        if interval_source == "drilling":
            return is_facies and is_drilling
        if interval_source == "gis":
            return is_facies and is_gis and not is_drilling
        return True

    tops = []
    bases = []
    for index, value in enumerate(headers):
        column = index + 1
        if column in used or not belongs_to_source(value):
            continue
        top_score = _score(value, "top")
        base_score = _score(value, "base")
        if top_score > 0:
            tops.append((column, top_score))
        if base_score > 0:
            bases.append((column, base_score))
    candidates: list[tuple[float, int, int]] = []
    for top_column, top_score in tops:
        for base_column, base_score in bases:
            if top_column == base_column:
                continue
            valid_pairs = 0
            thickness_checks = 0
            thickness_matches = 0
            for row in rows[header_row:]:
                top = as_float(_cell(row, top_column))
                base = as_float(_cell(row, base_column))
                if top is None or base is None or base <= top:
                    continue
                valid_pairs += 1
                thickness = as_float(_cell(row, thickness_column))
                if thickness is None:
                    continue
                thickness_checks += 1
                span_cm = meters_to_centimeters(base) - meters_to_centimeters(top)
                if abs(meters_to_centimeters(thickness) - span_cm) <= 1:
                    thickness_matches += 1
            score = float(top_score + base_score + min(valid_pairs, 20))
            if base_column == top_column + 1:
                score += 25
            if thickness_checks:
                match_ratio = thickness_matches / thickness_checks
                score += 500 * match_ratio - 400 * (1.0 - match_ratio)
            candidates.append((score, top_column, base_column))
    if not candidates:
        return None, None
    _, top_column, base_column = max(candidates, key=lambda item: (item[0], -item[1], -item[2]))
    return top_column, base_column


def _best_source_interval(headers: list[str], used: set[int], source: str) -> int | None:
    candidates = []
    for index, header in enumerate(headers, start=1):
        text = normalize_text(header)
        if index in used or not any(word in text for word in ("фаци", "слоя", "facies", "layer")):
            continue
        if any(word in text for word in ("кровл", "подошв", "верх", "низ", "top", "base", "bottom")):
            continue
        if source == "drilling" and not any(word in text for word in ("по бурен", "буров", "drilling")):
            continue
        if source == "gis" and (
            not any(word in text for word in ("по гис", "gis", "logging"))
            or any(word in text for word in ("по бурен", "буров", "drilling"))
        ):
            continue
        score = _score(header, "interval")
        if score:
            candidates.append((score, index))
    return max(candidates, default=(0, None))[1]


def detect_mapping(sheet: str, rows: list[list[Any]]) -> ColumnMapping:
    """Detect a multi-row header and semantic columns without a fixed template."""
    if not rows:
        return ColumnMapping(sheet=sheet, header_row=1)
    header_row = 1
    first_header_row = 1
    header_started = False
    for index, row in enumerate(rows[:30], start=1):
        nonempty = [normalize_text(value) for value in row if display_text(value)]
        if len(nonempty) > 1 and len(set(nonempty)) == 1:
            # A merged document title copied across the sheet must not add
            # "керн" or "фация" to every otherwise unrelated column.
            continue
        hits = {
            role
            for value in row
            for role in ROLE_ALIASES
            if _score(display_text(value), role) >= 100
        }
        structural = hits & {"well", "interval", "top", "base", "core_top", "core_base", "label", "class_code", "class_index"}
        starts_header = len(structural) >= 2 or (
            bool(structural & {"interval", "top", "base", "core_top", "core_base"})
            and bool(structural & {"well", "label", "class_code", "class_index"})
        )
        continues_header = header_started and index <= header_row + 2 and len(structural) >= 2
        if starts_header or continues_header:
            if not header_started:
                first_header_row = index
            header_row = index
            header_started = True
    width = max((len(row) for row in rows[:header_row]), default=0)
    headers = []
    for column in range(width):
        parts: list[str] = []
        for row in rows[first_header_row - 1:header_row]:
            value = display_text(row[column]) if column < len(row) else ""
            if value and (not parts or normalize_text(parts[-1]) != normalize_text(value)):
                parts.append(value)
        headers.append(" ".join(parts))

    used: set[int] = set()
    values: dict[str, int | None] = {}
    # Specific core fields must claim their columns before generic depth fields.
    for role in (
        "core_top", "core_base", "well", "class_code", "class_index", "label",
        "target_text", "association", "environment", "description", "field_name",
        "facies_thickness",
    ):
        values[role] = _best(headers, role, used)
        if values[role]:
            used.add(int(values[role]))
    values["top"], values["base"] = _best_facies_pair(
        headers, rows, header_row, used, values.get("facies_thickness"), "drilling",
    )
    drilling_interval = _best_source_interval(headers, used, "drilling")
    # Some workbooks contain only GIS facies limits. Use those as the primary
    # limits only when no drilling limits exist; otherwise retain both systems
    # so matching can use GIS strictly as a fallback.
    if not (values["top"] and values["base"]) and not drilling_interval:
        values["top"], values["base"] = _best_facies_pair(
            headers, rows, header_row, used, values.get("facies_thickness"),
        )
    if values["top"]:
        used.add(int(values["top"]))
    if values["base"]:
        used.add(int(values["base"]))
    values["interval"] = None
    if not (values["top"] and values["base"]):
        values["interval"] = drilling_interval
        if not values["interval"]:
            values["interval"] = _best(headers, "interval", used)
    values["gis_top"], values["gis_base"] = _best_facies_pair(
        headers, rows, header_row, used, values.get("facies_thickness"), "gis",
    )
    if not (values["gis_top"] and values["gis_base"]):
        values["gis_top"] = values["gis_base"] = None
    if values["gis_top"] and values["gis_base"]:
        values["gis_interval"] = None
    elif values["interval"] and _is_gis_header(headers[values["interval"] - 1]):
        values["gis_interval"] = values["interval"]
    else:
        values["gis_interval"] = _best_source_interval(headers, used, "gis")
    return ColumnMapping(sheet=sheet, header_row=header_row, **values)


def read_table(path: Path, mapping_file: Path | None = None) -> tuple[list[DescriptionRow], list[ColumnMapping], list[Issue]]:
    path = Path(path).expanduser().resolve(strict=True)
    if path.suffix.lower() not in SUPPORTED_TABLES:
        raise ValueError(f"Неподдерживаемый формат таблицы: {path.suffix}")
    sheets = _read_sheets(path)
    overrides = _load_mapping_overrides(mapping_file)
    mappings: list[ColumnMapping] = []
    descriptions: list[DescriptionRow] = []
    issues: list[Issue] = []
    for sheet_name, rows in sheets:
        mapping_key = f"{path}::{sheet_name}"
        detected = detect_mapping(sheet_name, rows)
        override = overrides.get(mapping_key) or overrides.get(sheet_name)
        if override is None:
            mapping = detected
        else:
            # Saved column maps from older versions intentionally remain
            # authoritative for the primary drilling columns. New semantic
            # fallback fields (GIS) are auto-discovered when absent.
            override_gis_pair = bool(override.gis_top and override.gis_base)
            detected_gis_pair = bool(detected.gis_top and detected.gis_base)
            if override_gis_pair:
                gis_top, gis_base = override.gis_top, override.gis_base
            elif detected_gis_pair:
                gis_top, gis_base = detected.gis_top, detected.gis_base
            else:
                gis_top, gis_base = override.gis_top, override.gis_base
            gis_interval = None if gis_top and gis_base else (override.gis_interval or detected.gis_interval)
            mapping = replace(
                override,
                gis_interval=gis_interval,
                gis_top=gis_top,
                gis_base=gis_base,
            )
        mapping = replace(mapping, source_file=str(path))
        mappings.append(mapping)
        parsed, sheet_issues = _parse_sheet(rows, mapping)
        descriptions.extend(parsed)
        issues.extend(sheet_issues)
    if not descriptions:
        raise ValueError(
            "Не найдены строки с корректными интервалами и метками. Проверьте column_mapping.json: "
            "нужны well, top/base или interval, а также label либо code/index."
        )
    descriptions.sort(key=lambda item: (well_key(item.well), item.top, item.base, item.source_id))
    return descriptions, mappings, issues


def discover_table_files(inputs: Path | Iterable[Path]) -> list[Path]:
    values = [inputs] if isinstance(inputs, Path) else list(inputs)
    found: dict[str, Path] = {}
    for raw in values:
        path = Path(raw).expanduser().resolve(strict=True)
        candidates = path.rglob("*") if path.is_dir() else (path,)
        for candidate in candidates:
            if candidate.is_file() and candidate.suffix.lower() in SUPPORTED_TABLES:
                found[str(candidate).casefold()] = candidate
    return sorted(found.values(), key=lambda item: str(item).casefold())


def read_many_tables(
    inputs: Path | Iterable[Path], mapping_file: Path | None = None,
) -> tuple[list[DescriptionRow], list[ColumnMapping], list[Issue], list[Path]]:
    files = discover_table_files(inputs)
    if not files:
        raise ValueError("Не найдены Excel/CSV-файлы.")
    all_rows: list[DescriptionRow] = []
    all_mappings: list[ColumnMapping] = []
    all_issues: list[Issue] = []
    for path in files:
        try:
            rows, mappings, issues = read_table(path, mapping_file)
        except (OSError, ValueError, RuntimeError) as exc:
            all_issues.append(Issue("error", str(path), str(exc)))
            continue
        all_rows.extend(rows)
        all_mappings.extend(mappings)
        all_issues.extend(issues)
    if not all_rows:
        raise ValueError("Ни в одном файле не найдены пригодные строки описания.")
    all_rows.sort(key=lambda item: (well_key(item.well), item.top, item.base, item.source_id))
    return all_rows, all_mappings, all_issues, files


def save_mappings(path: Path, mappings: Iterable[ColumnMapping]) -> None:
    payload = {
        (f"{item.source_file}::{item.sheet}" if item.source_file else item.sheet): item.to_dict()
        for item in mappings
    }
    Path(path).write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _load_mapping_overrides(path: Path | None) -> dict[str, ColumnMapping]:
    if path is None or not Path(path).is_file():
        return {}
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    result = {}
    for key, value in payload.items():
        result[key] = ColumnMapping.from_dict(value)
    return result


def _read_sheets(path: Path) -> list[tuple[str, list[list[Any]]]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else _sniff_delimiter(path)
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            return [(path.stem, [list(row) for row in csv.reader(source, delimiter=delimiter)])]
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("Для старого формата .xls установите xlrd>=2.0.") from exc
        book = xlrd.open_workbook(path)
        return [(sheet.name, [sheet.row_values(row) for row in range(sheet.nrows)]) for sheet in book.sheets()]
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Для Excel установите openpyxl>=3.1.") from exc
    book = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    result: list[tuple[str, list[list[Any]]]] = []
    try:
        for sheet in book.worksheets:
            merged_values: dict[tuple[int, int], Any] = {}
            for merged in sheet.merged_cells.ranges:
                value = sheet.cell(merged.min_row, merged.min_col).value
                for row_index in range(merged.min_row, merged.max_row + 1):
                    for column_index in range(merged.min_col, merged.max_col + 1):
                        merged_values[(row_index, column_index)] = value
            rows: list[list[Any]] = []
            for row_index, source_row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [
                    value if value is not None else merged_values.get((row_index, column_index))
                    for column_index, value in enumerate(source_row, start=1)
                ]
                while values and values[-1] is None:
                    values.pop()
                rows.append(values)
            while rows and not any(value is not None for value in rows[-1]):
                rows.pop()
            result.append((sheet.title, rows))
    finally:
        book.close()
    return result


def _sniff_delimiter(path: Path) -> str:
    sample = path.read_text(encoding="utf-8-sig", errors="replace")[:8192]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        return ";"


def _cell(row: list[Any], column: int | None) -> Any:
    return row[column - 1] if column and column <= len(row) else None


def _is_gis_header(value: str) -> bool:
    text = normalize_text(value)
    return any(word in text for word in ("по гис", "gis", "logging")) and not any(
        word in text for word in ("по бурен", "буров", "drilling")
    )


def _parse_sheet(rows: list[list[Any]], mapping: ColumnMapping) -> tuple[list[DescriptionRow], list[Issue]]:
    output: list[DescriptionRow] = []
    issues: list[Issue] = []
    last_well = ""
    last_field = ""
    last_core_top = last_core_base = None
    primary_headers = " ".join(
        display_text(_cell(row, column))
        for row in rows[:mapping.header_row]
        for column in (mapping.top, mapping.base, mapping.interval) if column
    )
    primary_is_gis = _is_gis_header(primary_headers)
    sheet_well = "" if normalize_text(mapping.sheet) in {"лист1", "sheet1", "sheet", "данные", "data"} else mapping.sheet
    if not ((mapping.top and mapping.base) or mapping.interval):
        issues.append(Issue("error", mapping.sheet, "Не найдены столбцы интервала слоя."))
        return output, issues
    for row_number, row in enumerate(rows[mapping.header_row :], start=mapping.header_row + 1):
        mapped_columns = [
            value for key, value in mapping.to_dict().items()
            if key not in {"header_row", "sheet", "source_file"} and isinstance(value, int)
        ]
        numbered = [column for column in mapped_columns if display_text(_cell(row, column)) == str(column)]
        if len(numbered) >= 4:
            continue
        top = as_float(_cell(row, mapping.top))
        base = as_float(_cell(row, mapping.base))
        if top is None or base is None:
            top, base = parse_interval(_cell(row, mapping.interval))
        gis_top = as_float(_cell(row, mapping.gis_top))
        gis_base = as_float(_cell(row, mapping.gis_base))
        if mapping.gis_interval:
            gis_top, gis_base = parse_interval(_cell(row, mapping.gis_interval))
        if gis_top is None or gis_base is None or gis_base <= gis_top:
            gis_top = gis_base = None
        row_source = "gis" if primary_is_gis else "drilling"
        if (top is None or base is None or base <= top) and gis_top is not None and gis_base is not None:
            top, base = gis_top, gis_base
            row_source = "gis"
        if top is None or base is None or base <= top:
            continue
        top = normalize_depth(top)
        base = normalize_depth(base)
        if base <= top:
            continue
        thickness = as_float(_cell(row, mapping.facies_thickness))
        thickness_declared = thickness is not None
        if thickness is not None:
            thickness = normalize_depth(thickness)
        if gis_top == top and gis_base == base:
            # If there was no drilling pair, GIS already became the primary
            # pair above and must not be counted twice as a fallback source.
            gis_top = gis_base = None
        gis_thickness_valid = False
        if thickness is not None and gis_top is not None and gis_base is not None:
            gis_span_cm = meters_to_centimeters(gis_base) - meters_to_centimeters(gis_top)
            gis_thickness_valid = abs(meters_to_centimeters(thickness) - gis_span_cm) <= 1
        thickness_valid = True
        if thickness is not None:
            span_cm = meters_to_centimeters(base) - meters_to_centimeters(top)
            thickness_valid = abs(meters_to_centimeters(thickness) - span_cm) <= 1
            if not thickness_valid:
                issues.append(Issue(
                    "warning" if gis_thickness_valid else "error",
                    f"{mapping.sheet}!{row_number}",
                    (
                        f"Толщина фации {format_depth(thickness)} м не совпадает с интервалом по бурению "
                        f"{format_depth(top)}–{format_depth(base)} м "
                        f"({format_depth(span_cm / 100)} м); резервный интервал ГИС совпадает с толщиной "
                        "и может применяться при сопоставлении фото в системе глубин ГИС."
                        if gis_thickness_valid else
                        f"Толщина фации {format_depth(thickness)} м не совпадает с интервалом фации по бурению "
                        f"{format_depth(top)}–{format_depth(base)} м "
                        f"({format_depth(span_cm / 100)} м); строка не будет использована для маски."
                    ),
                ))
        well = display_text(_cell(row, mapping.well)) or last_well or sheet_well
        if not well:
            issues.append(Issue("error", f"{mapping.sheet}!{row_number}", "Не указана скважина."))
            continue
        if last_well and well_key(well) != well_key(last_well):
            # Forward-fill within one well only. Otherwise the first row of a
            # new well inherits the previous well's sampling interval/field.
            last_core_top = last_core_base = None
            last_field = ""
        last_well = well
        field_name = display_text(_cell(row, mapping.field_name)) or last_field
        last_field = field_name
        core_top = as_float(_cell(row, mapping.core_top))
        core_base = as_float(_cell(row, mapping.core_base))
        if core_top is None:
            core_top = last_core_top
        else:
            last_core_top = core_top
        if core_base is None:
            core_base = last_core_base
        else:
            core_base = normalize_depth(core_base)
            last_core_base = core_base
        if core_top is not None:
            core_top = normalize_depth(core_top)
            last_core_top = core_top
        name = display_text(_cell(row, mapping.label))
        code = display_text(_cell(row, mapping.class_code))
        index = display_text(_cell(row, mapping.class_index))
        if code and index:
            label = f"{code}@{index}"
        else:
            label = code or index or name
        if not label:
            issues.append(Issue("error", f"{mapping.sheet}!{row_number}", "Нет метки класса; строка пропущена."))
            continue
        description = display_text(_cell(row, mapping.description))
        target_text = display_text(_cell(row, mapping.target_text)) or description
        association = display_text(_cell(row, mapping.association))
        environment = display_text(_cell(row, mapping.environment))
        metadata = {
            "field": field_name,
            "well": well,
            "description": description,
            "target_text": target_text,
            "association": association,
            "environment": environment,
            "source": f"{mapping.sheet}!{row_number}",
            "interval_source": row_source,
        }
        output.append(DescriptionRow(
            well=well, top=top, base=base, label=label, sheet=mapping.sheet, row=row_number,
            description=description, target_text=target_text, association=association,
            environment=environment, field_name=field_name, source_file=mapping.source_file,
            core_top=core_top, core_base=core_base,
            thickness=thickness if thickness is not None else normalize_depth(base - top),
            thickness_valid=thickness_valid,
            metadata={key: value for key, value in metadata.items() if value},
            gis_top=gis_top, gis_base=gis_base,
            thickness_declared=thickness_declared,
        ))
    return output, issues
