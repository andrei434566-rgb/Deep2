from __future__ import annotations

from pathlib import Path
from typing import Iterable, Mapping


COLUMNS = (
    ("field_name", "Месторождение", None),
    ("well", "№ скв.", None),
    ("core_run", "№ долбления", None),
    ("core_top", "Интервал отбора\nкерна, м", "Кровля"),
    ("core_base", "", "Подошва"),
    ("stratigraphy", "Стратиграфия", None),
    ("gis_shift", "Смещение по\nГИС, м", None),
    ("gis_core_top", "Интервал отбора\nпо ГИС, м", "Кровля"),
    ("gis_core_base", "", "Подошва"),
    ("drilled", "Проходка, м", None),
    ("recovered", "Вынос\nкерна, м", None),
    ("recovery_percent", "Вынос %", None),
    ("facies_top", "Интервал фации по бурению, м", "Кровля"),
    ("facies_base", "", "Подошва"),
    ("gis_facies_top", "Интервал фации по ГИС, м", "Кровля"),
    ("gis_facies_base", "", "Подошва"),
    ("layer_no", "№ слоя", None),
    ("thickness", "Толщина фации, м", None),
    ("facies_name", "Название фации", None),
    ("association", "Ассоциация фаций\n(по гидродинамическому режиму осадконакопления)", None),
    ("environment", "Обстановка осадконакопления", None),
    ("description", "Краткое описание", None),
)


def export_standardized_workbook(rows: Iterable[Mapping], destination: Path) -> Path:
    """Write the stable 22-column exchange format consumed by Kern Analyzer."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Для экспорта Excel установите openpyxl>=3.1.") from exc
    destination = Path(destination).expanduser().absolute()
    if destination.exists():
        raise FileExistsError(f"Файл уже существует: {destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Описание"
    for column, (_, title, subtitle) in enumerate(COLUMNS, start=1):
        sheet.cell(1, column, title)
        if subtitle:
            sheet.cell(2, column, subtitle)
        else:
            sheet.merge_cells(start_row=1, start_column=column, end_row=2, end_column=column)
        sheet.cell(3, column, column)
    for first, second in ((4, 5), (8, 9), (13, 14), (15, 16)):
        sheet.merge_cells(start_row=1, start_column=first, end_row=1, end_column=second)
    thin = Side(style="thin", color="333333")
    header_fill = PatternFill("solid", fgColor="D9EAF2")
    for row in sheet.iter_rows(min_row=1, max_row=3, min_col=1, max_col=22):
        for cell in row:
            cell.font = Font(name="Arial", size=9, bold=cell.row <= 2)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = header_fill
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row_index, values in enumerate(rows, start=4):
        for column, (key, _, _) in enumerate(COLUMNS, start=1):
            value = values.get(key)
            if key == "thickness" and value is None:
                top, base = values.get("facies_top"), values.get("facies_base")
                value = round(float(base) - float(top), 4) if top is not None and base is not None else None
            cell = sheet.cell(row_index, column, value)
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=column >= 19)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if column in {4, 5, 7, 8, 9, 10, 11, 13, 14, 15, 16, 18}:
                cell.number_format = "0.00"
        sheet.row_dimensions[row_index].height = 54
    widths = (18, 12, 11, 12, 12, 14, 12, 12, 12, 11, 11, 10, 13, 13, 13, 13, 9, 12, 18, 32, 27, 55)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 58
    sheet.row_dimensions[2].height = 24
    sheet.row_dimensions[3].height = 20
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:V{max(3, sheet.max_row)}"
    workbook.save(destination)
    return destination
