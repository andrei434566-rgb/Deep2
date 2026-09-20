from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from excel_photo_model_studio.standard_excel import export_standardized_workbook


class StandardExcelTests(unittest.TestCase):
    def test_exports_22_columns_with_description_in_column_22(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "well.xlsx"
            export_standardized_workbook([{
                "field_name": "Уренгойское", "well": "Р-31", "core_top": 3183.0,
                "core_base": 3195.0, "facies_top": 3188.98, "facies_base": 3189.50,
                "facies_name": "Tcr", "description": "Песчаник светло-серый, слоистый.",
            }], output)
            workbook = load_workbook(output, data_only=True)
            sheet = workbook.active
        self.assertEqual(22, sheet.max_column)
        self.assertEqual("Краткое описание", sheet.cell(1, 22).value)
        self.assertEqual("Песчаник светло-серый, слоистый.", sheet.cell(4, 22).value)
        self.assertAlmostEqual(0.52, sheet.cell(4, 18).value)


if __name__ == "__main__":
    unittest.main()
