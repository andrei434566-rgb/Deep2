from __future__ import annotations

import json
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from openpyxl import Workbook

from excel_photo_model_studio.depth import meters_to_centimeters
from excel_photo_model_studio.tabular import as_float, parse_interval, read_table, save_mappings


class TableReaderTests(unittest.TestCase):
    def test_excel_depths_accept_comma_dot_and_grouped_decimal_formats(self):
        expected = 4105.25
        for value in (4105.25, "4105.25", "4105,25", "4.105,25", "4,105.25", "4 105,25"):
            with self.subTest(value=value):
                self.assertEqual(expected, as_float(value))
                self.assertEqual(410525, meters_to_centimeters(value))
        self.assertEqual(0.96, as_float(".96"))
        self.assertEqual(0.96, as_float(",96"))

    def test_interval_parser_does_not_split_decimal_comma_or_dot(self):
        for value in ("4105.00-4108.65", "4105,00–4108,65", "4.105,00-4.108,65"):
            with self.subTest(value=value):
                self.assertEqual((4105.0, 4108.65), parse_interval(value))

    def test_missing_drilling_cells_do_not_drop_valid_gis_facies(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "missing_drilling.csv"
            path.write_text(
                "Скважина;Интервал фации по бурению Кровля;Интервал фации по бурению Подошва;"
                "Интервал фации по ГИС Кровля;Интервал фации по ГИС Подошва;"
                "Толщина фации;Код фации;Краткое описание\n"
                "W-1;;;200;201;1;A;Песчаник.\n", encoding="utf-8",
            )
            rows, _, issues = read_table(path)

        self.assertEqual(1, len(rows))
        self.assertEqual((200, 201), (rows[0].top, rows[0].base))
        self.assertEqual("gis", rows[0].metadata["interval_source"])
        self.assertEqual("Песчаник.", rows[0].target_text)
        self.assertFalse(any(item.severity == "error" for item in issues))

    def test_combined_drilling_interval_is_not_displaced_by_gis_pair(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "combined.csv"
            path.write_text(
                "Скважина;Интервал фации по бурению;Интервал фации по ГИС Кровля;"
                "Интервал фации по ГИС Подошва;Толщина фации;Код фации\n"
                "W-1;100-101;200;201;1;A\n", encoding="utf-8",
            )
            rows, mappings, _ = read_table(path)

        self.assertEqual(2, mappings[0].interval)
        self.assertEqual((100, 101), (rows[0].top, rows[0].base))
        self.assertEqual((200, 201), (rows[0].gis_top, rows[0].gis_base))

    def test_merged_title_does_not_turn_facies_pair_into_core_sampling_pair(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "reordered.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Послойное описание керна"])
            sheet.merge_cells("A1:H1")
            sheet.append([
                "Скважина", "Интервал фации по бурению Кровля", "Интервал фации по бурению Подошва",
                "Код фации", "Интервал отбора керна Кровля", "Интервал отбора керна Подошва",
                "Толщина фации", "Краткое описание",
            ])
            sheet.append(["W-1", 100, 101, "A", 100, 110, 1, "Песчаник."])
            sheet.append(["W-2", 200, 201, "B", None, None, 1, "Алевролит."])
            workbook.save(path)
            rows, mappings, _ = read_table(path)

        self.assertEqual((5, 6), (mappings[0].core_top, mappings[0].core_base))
        self.assertEqual((2, 3), (mappings[0].top, mappings[0].base))
        self.assertIsNone(rows[1].core_top)
        self.assertIsNone(rows[1].core_base)

    def test_reads_reference_22_column_schema_and_target_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "reference.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            headers = [
                "Месторождение", "№ скв.", "№ долбления", "Интервал отбора керна, м", "",
                "Стратиграфия", "Смещение по ГИС, м", "Интервал отбора по ГИС, м", "",
                "Проходка, м", "Вынос керна, м", "Вынос %", "Интервал фации по бурению, м", "",
                "Интервал фации по ГИС, м", "", "№ слоя", "Толщина фации, м",
                "Название фации", "Ассоциация фаций", "Обстановка осадконакопления", "Краткое описание",
            ]
            sheet.append(headers)
            for start, end in ((4, 5), (8, 9), (13, 14), (15, 16)):
                sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
            sheet.append([None, None, None, "Кровля", "Подошва", None, None, "Кровля", "Подошва", None, None, None, "Кровля", "Подошва", "Кровля", "Подошва"])
            sheet.append(list(range(1, 23)))
            sheet.append([
                "Уренгойское", "Р-31", 1, 3183.0, 3195.0, "К2m", -1, 3182.0, 3194.0,
                3, 3, 100, 3188.98, 3189.50, 3187.98, 3188.50, 1, 0.52,
                "Tcr", "Прибрежная равнина", "Прибрежная равнина", "Песчаник светло-серый, слоистый.",
            ])
            workbook.save(path)

            rows, mappings, _ = read_table(path)

        self.assertEqual(1, len(rows))
        self.assertEqual((4, 5), (mappings[0].core_top, mappings[0].core_base))
        self.assertEqual((13, 14), (mappings[0].top, mappings[0].base))
        self.assertEqual(22, mappings[0].target_text)
        self.assertEqual("Tcr", rows[0].label)
        self.assertEqual("Песчаник светло-серый, слоистый.", rows[0].target_text)

    def test_reads_merged_multiline_headers_and_forward_fills_well(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "description.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Описание"
            sheet["A1"] = "№ скважины"
            sheet["B1"] = "Интервал фации"
            sheet.merge_cells("B1:C1")
            sheet["B2"] = "Кровля"
            sheet["C2"] = "Подошва"
            sheet["D1"] = "Код фации"
            sheet["E1"] = "Индекс фации"
            sheet.append(["Р-31", 3915.0, 3915.5, "Dch", "47"])
            sheet.append([None, 3915.5, 3916.0, "DWCh", "80"])
            workbook.save(path)

            rows, mappings, issues = read_table(path)

        self.assertEqual([], [item for item in issues if item.severity == "error"])
        self.assertEqual(2, len(rows))
        self.assertEqual("Р-31", rows[1].well)
        self.assertEqual("DWCh@80", rows[1].label)
        self.assertEqual(2, mappings[0].header_row)

    def test_finds_drilling_interval_and_short_description_by_header_not_position(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "extended_31_columns.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "седимент"
            sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=31)
            sheet.cell(1, 1, "Послойное седиментологическое описание керна")
            headers = [
                "Месторождение", "№ скв.", "№ долбления", "Интервал отбора керна, м", None,
                "Стратиграфия", "Смещение по ГИС, м", "Интервал отбора по ГИС, м", None,
                "Проходка, м", "Вынос керна, м", "Вынос, %", "№", "Интервал фации по бурению, м", None,
                "Интервал фации по ГИС, м", None, "№ слоя", "Толщина фации, м", "Индекс фации",
                "Название фации", "Код фации", "Индекс макрофации", "Название макрофации",
                "Код макрофации", "Индекс ассоциации фаций", "Ассоциация фаций",
                "Обстановка осадконакопления", "Комплекс осадконакопления", "Краткое описание", "Примечание",
            ]
            sheet.append(headers)
            for start, end in ((4, 5), (8, 9), (14, 15), (16, 17)):
                sheet.merge_cells(start_row=2, start_column=start, end_row=2, end_column=end)
            subheaders = [None] * 31
            for start, end in ((4, 5), (8, 9), (14, 15), (16, 17)):
                subheaders[start - 1] = "Кровля"
                subheaders[end - 1] = "Подошва"
            sheet.append(subheaders)
            sheet.append(list(range(1, 32)))
            data = [None] * 31
            values = {
                1: "Восточно-Тазовское", 2: "67ПО", 4: 4105.0, 5: 4117.0,
                14: 4105.0, 15: 4108.65, 16: 4104.90, 17: 4108.55,
                19: 3.65, 20: "Dch", 21: "Каналы распределительные", 22: 92,
                27: "Дельтовая ассоциация", 28: "Дельтовая обстановка",
                30: "Песчаник светло-серый, слоистый.", 31: "Контроль",
            }
            for column, value in values.items():
                data[column - 1] = value
            sheet.append(data)
            workbook.save(path)

            rows, mappings, issues = read_table(path)

        self.assertEqual([], [item for item in issues if item.severity == "error"])
        self.assertEqual((14, 15), (mappings[0].top, mappings[0].base))
        self.assertEqual((16, 17), (mappings[0].gis_top, mappings[0].gis_base))
        self.assertEqual(19, mappings[0].facies_thickness)
        self.assertEqual(30, mappings[0].target_text)
        self.assertEqual("Песчаник светло-серый, слоистый.", rows[0].target_text)
        self.assertEqual((4105.0, 4108.65), (rows[0].top, rows[0].base))
        self.assertEqual((4104.9, 4108.55), (rows[0].gis_top, rows[0].gis_base))
        self.assertTrue(rows[0].thickness_valid)

    def test_thickness_mismatch_blocks_row_from_masking(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bad_thickness.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append([
                "Скважина", "Интервал фации по бурению Кровля",
                "Интервал фации по бурению Подошва", "Толщина фации, м", "Класс",
            ])
            sheet.append(["W-1", 100.0, 112.0, 3.65, "Dch"])
            workbook.save(path)

            rows, mappings, issues = read_table(path)

        self.assertEqual((2, 3, 4), (mappings[0].top, mappings[0].base, mappings[0].facies_thickness))
        self.assertEqual(1, len(rows))
        self.assertFalse(rows[0].thickness_valid)
        self.assertTrue(any("не будет использована для маски" in item.message for item in issues))

    def test_rounds_formula_noise_to_centimetres_and_treats_3_03_as_303_cm(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "centimetres.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append([
                "Скважина", "Интервал фации по бурению Кровля",
                "Интервал фации по бурению Подошва", "Толщина фации, м", "Класс",
            ])
            sheet.append(["W-1", 4129.67, 4130.139999999, 0.47, "A"])
            sheet.append(["W-1", 4130.139999999, 4133.17, 3.03, "B"])
            workbook.save(path)

            rows, _, issues = read_table(path)

        self.assertEqual([(4129.67, 4130.14), (4130.14, 4133.17)], [
            (row.top, row.base) for row in rows
        ])
        self.assertEqual([0.47, 3.03], [row.thickness for row in rows])
        self.assertTrue(all(row.thickness_valid for row in rows))
        self.assertEqual([], [item for item in issues if item.severity == "error"])

    def test_reads_csv_with_one_interval_column(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "description.csv"
            path.write_text("Скважина;Интервал;Класс\nW-1;100,0-101,5;Sand\n", encoding="utf-8")

            rows, _, _ = read_table(path)

        self.assertEqual(1, len(rows))
        self.assertEqual((100.0, 101.5), (rows[0].top, rows[0].base))
        self.assertEqual("Sand", rows[0].label)

    def test_reads_gis_facies_limits_as_a_separate_fallback(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "drilling_and_gis.csv"
            path.write_text(
                "Скважина;Интервал фации по бурению Кровля;Интервал фации по бурению Подошва;"
                "Интервал фации по ГИС Кровля;Интервал фации по ГИС Подошва;"
                "Толщина фации;Код фации;Краткое описание\n"
                "W-1;100.00;101.00;200.00;201.00;1.00;Sand;Песчаник серый.\n",
                encoding="utf-8",
            )

            rows, mappings, issues = read_table(path)

        self.assertEqual([], [item for item in issues if item.severity == "error"])
        self.assertEqual((2, 3), (mappings[0].top, mappings[0].base))
        self.assertEqual((4, 5), (mappings[0].gis_top, mappings[0].gis_base))
        self.assertEqual((100.0, 101.0), (rows[0].top, rows[0].base))
        self.assertEqual((200.0, 201.0), (rows[0].gis_top, rows[0].gis_base))

    def test_valid_gis_thickness_turns_bad_drilling_interval_into_a_warning(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "gis_rescue.csv"
            path.write_text(
                "Скважина;Интервал фации по бурению Кровля;Интервал фации по бурению Подошва;"
                "Интервал фации по ГИС Кровля;Интервал фации по ГИС Подошва;"
                "Толщина фации;Код фации\nW-1;100.00;102.00;200.00;201.00;1.00;Sand\n",
                encoding="utf-8",
            )

            rows, _, issues = read_table(path)

        self.assertFalse(rows[0].thickness_valid)
        self.assertEqual((200.0, 201.0), (rows[0].gis_top, rows[0].gis_base))
        self.assertTrue(any(item.severity == "warning" and "резервный интервал ГИС" in item.message for item in issues))
        self.assertFalse(any(item.severity == "error" and "Толщина фации" in item.message for item in issues))

    def test_old_saved_column_map_is_supplemented_with_new_gis_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "old_mapping.csv"
            path.write_text(
                "Скважина;Интервал фации по бурению Кровля;Интервал фации по бурению Подошва;"
                "Интервал фации по ГИС Кровля;Интервал фации по ГИС Подошва;"
                "Толщина фации;Код фации\nW-1;100.00;101.00;200.00;201.00;1.00;Sand\n",
                encoding="utf-8",
            )
            _, detected, _ = read_table(path)
            mapping_file = Path(directory) / "column_mapping.json"
            save_mappings(mapping_file, [replace(detected[0], gis_top=None, gis_base=None)])
            old_map = json.loads(mapping_file.read_text(encoding="utf-8"))
            for saved in old_map.values():
                saved.pop("gis_interval", None)
                saved.pop("gis_top", None)
                saved.pop("gis_base", None)
            mapping_file.write_text(json.dumps(old_map), encoding="utf-8")

            rows, mappings, _ = read_table(path, mapping_file)

        self.assertEqual((2, 3), (mappings[0].top, mappings[0].base))
        self.assertEqual((4, 5), (mappings[0].gis_top, mappings[0].gis_base))
        self.assertEqual((200.0, 201.0), (rows[0].gis_top, rows[0].gis_base))

    def test_uses_gis_as_primary_interval_when_workbook_has_no_drilling_limits(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "gis_only.csv"
            path.write_text(
                "Скважина;Интервал фации по ГИС Кровля;Интервал фации по ГИС Подошва;"
                "Толщина фации;Код фации\nW-1;200.00;201.00;1.00;Sand\n",
                encoding="utf-8",
            )

            rows, mappings, issues = read_table(path)

        self.assertEqual([], [item for item in issues if item.severity == "error"])
        self.assertEqual((2, 3), (mappings[0].top, mappings[0].base))
        self.assertEqual((200.0, 201.0), (rows[0].top, rows[0].base))


if __name__ == "__main__":
    unittest.main()
